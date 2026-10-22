#!/usr/bin/env python3
"""Add (or update) a team: logo -> assets/teams/<slug>.webp, color from logo -> data/teams.json,
and team name -> first empty cell in column C of the workbook (Table2).

usage: python3 tools/add_team.py "TEAM NAME" path/to/logo.(png|jpg|webp) [--no-excel]
"""
import sys, os, re, json, colorsys, shutil
import numpy as np
from PIL import Image
from scipy import ndimage

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEAMS = os.path.join(ROOT, "data", "teams.json")
XLSX = os.path.join(ROOT, "data", "DYNASTY GLORY LEAGUE.xlsx")
OUT = os.path.join(ROOT, "assets", "teams")
SIZE = 256


def slug(n):
    s = re.sub(r"[^a-z0-9]+", "-", n.lower()).strip("-")
    return s or "team"


def cut_background(im):
    """If the logo has no transparency, remove a flat background connected to the borders."""
    a = np.array(im.convert("RGBA")).astype(np.int16)
    if (a[..., 3] < 250).mean() > 0.02:
        return Image.fromarray(a.astype(np.uint8))
    rgb = a[..., :3]
    edge = np.concatenate([rgb[0], rgb[-1], rgb[:, 0], rgb[:, -1]])
    bg = np.median(edge, axis=0)
    if np.abs(edge - bg).max(axis=1).mean() > 40:  # busy border -> keep as is
        return Image.fromarray(a.astype(np.uint8))
    near = np.abs(rgb - bg).max(axis=2) < 38
    lab, _ = ndimage.label(near)
    border = set(np.unique(np.concatenate([lab[0], lab[-1], lab[:, 0], lab[:, -1]]))) - {0}
    mask = np.isin(lab, list(border))
    mask = ndimage.binary_dilation(mask, iterations=1) & near
    a[..., 3][mask] = 0
    return Image.fromarray(a.astype(np.uint8))


def square(im):
    bb = im.getchannel("A").point(lambda v: 255 if v > 12 else 0).getbbox()
    if bb:
        im = im.crop(bb)
    w, h = im.size
    side = int(max(w, h) * 1.04)
    c = Image.new("RGBA", (side, side), (0, 0, 0, 0))
    c.paste(im, ((side - w) // 2, (side - h) // 2), im)
    return c.resize((SIZE, SIZE), Image.LANCZOS)


def main_color(im):
    a = np.array(im.convert("RGBA")).reshape(-1, 4).astype(float) / 255
    a = a[a[:, 3] > 0.8][:, :3]
    if not len(a):
        return "#888888"
    hsv = np.array([colorsys.rgb_to_hsv(*p) for p in a[:: max(1, len(a) // 20000)]])
    rgb = a[:: max(1, len(a) // 20000)]
    ok = (hsv[:, 1] > 0.28) & (hsv[:, 2] > 0.22)
    if ok.sum() < len(hsv) * 0.04:  # almost monochrome logo
        v = rgb.mean(axis=0)
    else:
        h, s, rr = hsv[ok], hsv[ok, 1], rgb[ok]
        bins = (h[:, 0] * 18).astype(int) % 18
        score = np.bincount(bins, weights=s, minlength=18)
        b = score.argmax()
        sel = (bins == b) | (bins == (b + 1) % 18) | (bins == (b - 1) % 18)
        v = rr[sel].mean(axis=0)
    hh, ss, vv = colorsys.rgb_to_hsv(*v)
    vv = max(vv, 0.62)  # stay readable on the dark #090F1C background
    v = colorsys.hsv_to_rgb(hh, ss, vv)
    return "#%02X%02X%02X" % tuple(int(round(x * 255)) for x in v)


def add_excel(name):
    import openpyxl
    shutil.copy(XLSX, "/tmp/DYNASTY-before-add-team.xlsx")
    wb = openpyxl.load_workbook(XLSX)
    ws = wb["DYNASTY GLORY"]
    names = [ws[f"C{r}"].value for r in range(3, 63)]
    if name in names:
        return "already in workbook"
    for r in range(3, 63):
        if ws[f"C{r}"].value in (None, ""):
            ws[f"C{r}"].value = name
            wb.save(XLSX)
            chk = openpyxl.load_workbook(XLSX)["DYNASTY GLORY"][f"C{r}"].value
            assert chk == name
            return f"C{r}"
    return "no empty row"


def main():
    args = [x for x in sys.argv[1:] if not x.startswith("--")]
    if len(args) != 2:
        sys.exit(__doc__)
    name, src = args[0].strip(), args[1]
    os.makedirs(OUT, exist_ok=True)
    im = square(cut_background(Image.open(src)))
    fn = slug(name) + ".webp"
    im.save(os.path.join(OUT, fn), "WEBP", quality=92, method=6)
    color = main_color(im)
    teams = json.load(open(TEAMS, encoding="utf-8")) if os.path.exists(TEAMS) else []
    entry = {"name": name, "color": color, "logo": "assets/teams/" + fn}
    for i, t in enumerate(teams):
        if t["name"] == name:
            teams[i] = entry
            break
    else:
        teams.append(entry)
    json.dump(teams, open(TEAMS, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    xl = "skipped" if "--no-excel" in sys.argv else add_excel(name)
    print(json.dumps({"team": entry, "excel": xl, "count": len(teams)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
